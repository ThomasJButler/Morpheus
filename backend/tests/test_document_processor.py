import pytest

from app.utils.document_processor import (
    DocumentProcessingError,
    DocumentTooLarge,
    extract,
)
from tests.pdf_fixtures import make_pdf

CAPS = {"max_pdf_pages": 500, "max_text_chars": 2_000_000}


def test_txt(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Remote work is allowed three days a week.")
    out = extract(path, **CAPS)
    assert out == [{"text": "Remote work is allowed three days a week.", "page": None}]


def test_md(tmp_path):
    path = tmp_path / "readme.md"
    path.write_text("# Title\n\nSome body text.")
    out = extract(path, **CAPS)
    assert out[0]["text"].startswith("# Title")


def test_non_utf8_falls_back(tmp_path):
    path = tmp_path / "legacy.txt"
    path.write_bytes(b"caf\xe9 menu")  # latin-1 e-acute
    out = extract(path, **CAPS)
    assert "caf" in out[0]["text"] and "menu" in out[0]["text"]


def test_empty_file_rejected(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("   ")
    with pytest.raises(DocumentProcessingError):
        extract(path, **CAPS)


def test_char_cap(tmp_path):
    path = tmp_path / "big.txt"
    path.write_text("x" * 100)
    with pytest.raises(DocumentTooLarge):
        extract(path, max_pdf_pages=500, max_text_chars=10)


def test_unsupported_suffix(tmp_path):
    path = tmp_path / "archive.zip"
    path.write_bytes(b"PK\x03\x04")
    with pytest.raises(DocumentProcessingError):
        extract(path, **CAPS)


def test_pdf_pages_and_text(tmp_path):
    path = tmp_path / "doc.pdf"
    path.write_bytes(make_pdf(["Hello page one", "Second page here"]))
    out = extract(path, **CAPS)
    assert [entry["page"] for entry in out] == [1, 2]
    assert "Hello page one" in out[0]["text"]
    assert "Second page here" in out[1]["text"]


def test_pdf_page_cap(tmp_path):
    path = tmp_path / "long.pdf"
    path.write_bytes(make_pdf(["a", "b", "c"]))
    with pytest.raises(DocumentTooLarge):
        extract(path, max_pdf_pages=2, max_text_chars=2_000_000)


def test_malformed_pdf_gives_fixed_message(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.4\nnot really a pdf at all")
    with pytest.raises(DocumentProcessingError) as exc:
        extract(path, **CAPS)
    # Fixed message, no internals, no temp paths (SECURITY_REVIEW F8).
    assert str(exc.value) == "Could not read this PDF."


def test_docx_paragraphs_and_tables(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("Laptops must use full disk encryption.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Region"
    table.rows[0].cells[1].text = "Leeds"
    path = tmp_path / "policy.docx"
    document.save(str(path))
    out = extract(path, **CAPS)
    text = out[0]["text"]
    assert "full disk encryption" in text
    assert "Region | Leeds" in text
